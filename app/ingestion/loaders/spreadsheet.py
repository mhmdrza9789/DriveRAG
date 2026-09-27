from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


def parse_spreadsheet(path: str | Path) -> str:
    path = Path(path)

    workbook = load_workbook(
        filename=path,
        read_only=True,
        data_only=True,
    )

    sections: list[str] = []

    try:
        for worksheet in workbook.worksheets:
            rows: list[str] = []

            for row in worksheet.iter_rows(values_only=True):
                values = [
                    str(value).strip()
                    for value in row
                    if value is not None and str(value).strip()
                ]

                if values:
                    rows.append(" | ".join(values))

            if rows:
                sections.append(
                    f"Sheet: {worksheet.title}\n" + "\n".join(rows)
                )

        return "\n\n".join(sections)

    finally:
        workbook.close()
