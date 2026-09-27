# # app/ingestion/loaders/office.py

def parse_docx(file_path: str) -> str:
    from docx import Document

    doc = Document(file_path)
    parts = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)

    for table in doc.tables:
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    row_text.append(cell_text)
            if row_text:
                parts.append(" | ".join(row_text))

    return "\n".join(parts)


def parse_pptx(file_path: str) -> str:
    from pptx import Presentation

    prs = Presentation(file_path)
    parts = []

    for slide_index, slide in enumerate(prs.slides, start=1):
        slide_parts = []

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = shape.text.strip()
                if text:
                    slide_parts.append(text)

        if slide_parts:
            parts.append(f"\n--- Slide {slide_index} ---\n" + "\n".join(slide_parts))

    return "\n".join(parts)


def parse_office(file_path: str) -> str:
    ext = file_path.lower().rsplit(".", 1)[-1]

    if ext == "docx":
        return parse_docx(file_path)

    if ext == "pptx":
        return parse_pptx(file_path)

    raise ValueError(f"Unsupported office file type: {ext}")
